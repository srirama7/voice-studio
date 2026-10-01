"""
Batch 1 Voice Synthesis Test Script for Voice Studio.

Tests 20 voices (7 cloned + 13 gallery voices) against 5 document file formats:
- QA_REPORT.md (.md)
- Roblox_AI_Testing_MCP_Specification_v2.pdf (.pdf)
- Amogha_Bhat_Cover_Letter.txt (.txt)
- SYNOPSIS ANAGHA(2).docx (.docx)
- HEART FAILURE PPT(3)_v2_backup.pptx (.pptx)

Generates test_batch1_results.json with detailed verification metrics.
"""

import os
import sys
import time
import json
import wave
import traceback
import numpy as np
from pathlib import Path

import fitz  # PyMuPDF
import docx
import pptx

from config import default_config
from document_parser import DocumentParser
from job_engine import JobEngine, JobStatus

# Master configuration
BASE_DIR = Path(r"C:\Users\amogh\my-voice-studio")
SAMPLE_DOCS_DIR = BASE_DIR / "test_sample_docs"
RESULTS_JSON_PATH = BASE_DIR / "test_batch1_results.json"

# Original test files
ORIGINAL_FILES = {
    "md": r"C:\Users\amogh\Downloads\QA_REPORT.md",
    "pdf": r"C:\Users\amogh\Downloads\Roblox_AI_Testing_MCP_Specification_v2.pdf",
    "txt": r"C:\Users\amogh\Downloads\Amogha_Bhat_Cover_Letter.txt",
    "docx": r"C:\Users\amogh\Downloads\SYNOPSIS ANAGHA(2).docx",
    "pptx": r"C:\Users\amogh\Downloads\HEART FAILURE PPT(3)_v2_backup.pptx",
}

# 20 Voices for Batch 1
VOICES_BATCH1 = [
    {"id": "speaker_01", "type": "clone", "gallery_voice": "", "lang": "en-in"},
    {"id": "amogh", "type": "clone", "gallery_voice": "", "lang": "en-in"},
    {"id": "nethuman", "type": "clone", "gallery_voice": "", "lang": "en-in"},
    {"id": "netvoice2", "type": "clone", "gallery_voice": "", "lang": "en-in"},
    {"id": "speaker_demo", "type": "clone", "gallery_voice": "", "lang": "en-in"},
    {"id": "test_speaker_01", "type": "clone", "gallery_voice": "", "lang": "en-in"},
    {"id": "voice3", "type": "clone", "gallery_voice": "", "lang": "en-in"},
    {"id": "gallery:prabhat_in_m", "type": "gallery", "gallery_voice": "prabhat_in_m", "lang": "en-in"},
    {"id": "gallery:prabhat_deep_in_m", "type": "gallery", "gallery_voice": "prabhat_deep_in_m", "lang": "en-in"},
    {"id": "gallery:prabhat_young_in_m", "type": "gallery", "gallery_voice": "prabhat_young_in_m", "lang": "en-in"},
    {"id": "gallery:neerja_in_f", "type": "gallery", "gallery_voice": "neerja_in_f", "lang": "en-in"},
    {"id": "gallery:neerjaexp_in_f", "type": "gallery", "gallery_voice": "neerjaexp_in_f", "lang": "en-in"},
    {"id": "gallery:neerja_soft_in_f", "type": "gallery", "gallery_voice": "neerja_soft_in_f", "lang": "en-in"},
    {"id": "gallery:neerja_lively_in_f", "type": "gallery", "gallery_voice": "neerja_lively_in_f", "lang": "en-in"},
    {"id": "gallery:madhur_hi_m", "type": "gallery", "gallery_voice": "madhur_hi_m", "lang": "hi"},
    {"id": "gallery:madhur_deep_hi_m", "type": "gallery", "gallery_voice": "madhur_deep_hi_m", "lang": "hi"},
    {"id": "gallery:madhur_young_hi_m", "type": "gallery", "gallery_voice": "madhur_young_hi_m", "lang": "hi"},
    {"id": "gallery:swara_hi_f", "type": "gallery", "gallery_voice": "swara_hi_f", "lang": "hi"},
    {"id": "gallery:swara_soft_hi_f", "type": "gallery", "gallery_voice": "swara_soft_hi_f", "lang": "hi"},
    {"id": "gallery:swara_lively_hi_f", "type": "gallery", "gallery_voice": "swara_lively_hi_f", "lang": "hi"},
]


def prepare_sample_documents() -> dict:
    """Extract sample text from original documents and create valid 1-page/slide sample test files."""
    SAMPLE_DOCS_DIR.mkdir(parents=True, exist_ok=True)
    sample_files = {}

    print("--- Preparing Sample Test Documents ---")
    for ext, orig_path in ORIGINAL_FILES.items():
        if not Path(orig_path).exists():
            raise FileNotFoundError(f"Original test file missing: {orig_path}")
        
        parsed = DocumentParser.parse(orig_path)
        first_slide_text = parsed.slides[0].get_narrative_text() if parsed.slides else "Voice studio test sample text."
        # Clean up text snippet for clean synthesis
        first_slide_text = first_slide_text.replace("\ufffd", "").strip()
        if len(first_slide_text) > 300:
            first_slide_text = first_slide_text[:300] + "."

        sample_path = SAMPLE_DOCS_DIR / f"sample_{ext}_doc.{ext}"

        if ext == "md":
            sample_path.write_text(f"# Slide 1: Sample QA Report\n\n{first_slide_text}", encoding="utf-8")
        elif ext == "txt":
            sample_path.write_text(first_slide_text, encoding="utf-8")
        elif ext == "pdf":
            doc = fitz.open()
            page = doc.new_page()
            page.insert_text((50, 50), first_slide_text)
            doc.save(str(sample_path))
            doc.close()
        elif ext == "docx":
            doc = docx.Document()
            doc.add_heading("Sample Document Section", level=1)
            doc.add_paragraph(first_slide_text)
            doc.save(str(sample_path))
        elif ext == "pptx":
            prs = pptx.Presentation()
            blank_layout = prs.slide_layouts[6]
            slide = prs.slides.add_slide(blank_layout)
            tx_box = slide.shapes.add_textbox(pptx.util.Inches(1), pptx.util.Inches(1), pptx.util.Inches(6), pptx.util.Inches(2))
            tx_box.text_frame.text = first_slide_text
            prs.save(str(sample_path))

        print(f"Created sample [{ext.upper()}]: {sample_path} (Text length: {len(first_slide_text)} chars)")
        sample_files[ext] = str(sample_path)

    return sample_files


def verify_audio_file(audio_path: str) -> dict:
    """Verify audio file exists, load waveform, check duration > 0, peak amplitude, and RMS."""
    if not audio_path or not Path(audio_path).exists():
        return {
            "valid": False,
            "error": f"Audio file does not exist: {audio_path}",
            "duration_sec": 0.0,
            "sample_rate": 0,
            "num_samples": 0,
            "peak_amplitude": 0.0,
            "rms": 0.0,
        }

    try:
        with wave.open(audio_path, "rb") as wf:
            channels = wf.getnchannels()
            sample_width = wf.getsampwidth()
            sample_rate = wf.getframerate()
            n_frames = wf.getnframes()
            raw_bytes = wf.readframes(n_frames)

        if n_frames == 0 or len(raw_bytes) == 0:
            return {
                "valid": False,
                "error": "Audio file contains 0 frames",
                "duration_sec": 0.0,
                "sample_rate": sample_rate,
                "num_samples": 0,
                "peak_amplitude": 0.0,
                "rms": 0.0,
            }

        duration_sec = n_frames / float(sample_rate)
        
        # Convert bytes to float samples
        if sample_width == 2:
            audio_samples = np.frombuffer(raw_bytes, dtype=np.int16).astype(np.float32) / 32768.0
        elif sample_width == 4:
            audio_samples = np.frombuffer(raw_bytes, dtype=np.int32).astype(np.float32) / 2147483648.0
        else:
            audio_samples = np.frombuffer(raw_bytes, dtype=np.uint8).astype(np.float32) / 128.0 - 1.0

        peak_amp = float(np.max(np.abs(audio_samples))) if len(audio_samples) > 0 else 0.0
        rms_val = float(np.sqrt(np.mean(audio_samples ** 2))) if len(audio_samples) > 0 else 0.0

        if duration_sec <= 0.0:
            return {
                "valid": False,
                "error": f"Invalid audio duration: {duration_sec}s",
                "duration_sec": duration_sec,
                "sample_rate": sample_rate,
                "num_samples": len(audio_samples),
                "peak_amplitude": peak_amp,
                "rms": rms_val,
            }

        return {
            "valid": True,
            "error": None,
            "duration_sec": round(duration_sec, 3),
            "sample_rate": sample_rate,
            "num_samples": len(audio_samples),
            "peak_amplitude": round(peak_amp, 4),
            "rms": round(rms_val, 4),
        }
    except Exception as exc:
        return {
            "valid": False,
            "error": f"Audio inspection error: {exc}",
            "duration_sec": 0.0,
            "sample_rate": 0,
            "num_samples": 0,
            "peak_amplitude": 0.0,
            "rms": 0.0,
        }


def run_batch1_tests():
    start_time = time.time()
    print("==========================================================")
    print("   STARTING BATCH 1 VOICE SYNTHESIS TEST SUITE (20 VOICES)")
    print("==========================================================")

    sample_files = prepare_sample_documents()
    engine = JobEngine(config=default_config)

    test_results = []
    voice_stats = {v["id"]: {"passed": 0, "failed": 0, "total": 0} for v in VOICES_BATCH1}
    format_stats = {ext: {"passed": 0, "failed": 0, "total": 0} for ext in ORIGINAL_FILES.keys()}

    total_runs = len(VOICES_BATCH1) * len(ORIGINAL_FILES)
    current_run = 0

    for v_idx, voice_info in enumerate(VOICES_BATCH1, start=1):
        voice_id = voice_info["id"]
        voice_type = voice_info["type"]
        gallery_voice = voice_info["gallery_voice"]
        lang = voice_info["lang"]

        print(f"\n[{v_idx}/20] Testing Voice: '{voice_id}' (Type: {voice_type}, Lang: {lang})")

        for ext, doc_path in sample_files.items():
            current_run += 1
            clean_vid = voice_id.replace(":", "_").replace("-", "_")
            job_id = f"test_b1_v{v_idx:02d}_{clean_vid}_{ext}"
            orig_doc_path = ORIGINAL_FILES[ext]

            print(f"  ({current_run}/{total_runs}) Document format: .{ext.upper()} ... ", end="", flush=True)

            job_start = time.time()
            test_entry = {
                "test_index": current_run,
                "voice_id": voice_id,
                "voice_type": voice_type,
                "gallery_voice": gallery_voice,
                "language": lang,
                "doc_extension": ext,
                "original_file": orig_doc_path,
                "sample_file": doc_path,
                "job_id": job_id,
                "passed": False,
                "error_message": None,
                "execution_time_sec": 0.0,
                "audio_metrics": {},
            }

            try:
                manifest = engine.run_job(
                    job_id=job_id,
                    document_path=doc_path,
                    voice_id=voice_id,
                    text_mode="exact",
                    language=lang,
                    gallery_voice=gallery_voice,
                )

                job_elapsed = time.time() - job_start
                test_entry["execution_time_sec"] = round(job_elapsed, 3)

                if manifest.status != JobStatus.COMPLETED:
                    test_entry["error_message"] = f"Job status is {manifest.status.value}, expected COMPLETED"
                    print(f"FAILED (Status: {manifest.status.value})")
                else:
                    audio_metrics = verify_audio_file(manifest.mastered_audio_file)
                    test_entry["audio_metrics"] = audio_metrics
                    test_entry["mastered_audio_file"] = manifest.mastered_audio_file

                    if audio_metrics["valid"]:
                        test_entry["passed"] = True
                        print(f"PASSED ({audio_metrics['duration_sec']}s, SR={audio_metrics['sample_rate']}Hz, RMS={audio_metrics['rms']})")
                    else:
                        test_entry["error_message"] = audio_metrics["error"]
                        print(f"FAILED ({audio_metrics['error']})")

            except Exception as exc:
                job_elapsed = time.time() - job_start
                test_entry["execution_time_sec"] = round(job_elapsed, 3)
                test_entry["error_message"] = str(exc)
                test_entry["traceback"] = traceback.format_exc()
                print(f"ERROR ({exc})")

            # Update statistics
            voice_stats[voice_id]["total"] += 1
            format_stats[ext]["total"] += 1

            if test_entry["passed"]:
                voice_stats[voice_id]["passed"] += 1
                format_stats[ext]["passed"] += 1
            else:
                voice_stats[voice_id]["failed"] += 1
                format_stats[ext]["failed"] += 1

            test_results.append(test_entry)

    total_elapsed = time.time() - start_time
    total_passed = sum(1 for r in test_results if r["passed"])
    total_failed = total_runs - total_passed
    pass_rate = (total_passed / float(total_runs)) * 100.0 if total_runs > 0 else 0.0

    print("\n==========================================================")
    print("              BATCH 1 TEST EXECUTION SUMMARY              ")
    print("==========================================================")
    print(f"Total Tests Executed: {total_runs}")
    print(f"Passed:               {total_passed}")
    print(f"Failed:               {total_failed}")
    print(f"Pass Rate:            {pass_rate:.2f}%")
    print(f"Total Time:           {total_elapsed:.2f} seconds")
    print("==========================================================")

    # Save detailed JSON results
    output_payload = {
        "summary": {
            "batch_name": "Batch 1 (Voices 1-20)",
            "timestamp": time.strftime("%Y-%m-%dT%H:%M:%SZ"),
            "total_tests": total_runs,
            "passed_tests": total_passed,
            "failed_tests": total_failed,
            "pass_rate_percent": round(pass_rate, 2),
            "execution_time_sec": round(total_elapsed, 2),
        },
        "voice_summary": voice_stats,
        "format_summary": format_stats,
        "test_results": test_results,
    }

    with open(RESULTS_JSON_PATH, "w", encoding="utf-8") as f:
        json.dump(output_payload, f, indent=2)

    print(f"\nDetailed JSON report saved to: {RESULTS_JSON_PATH}")
    return output_payload


if __name__ == "__main__":
    run_batch1_tests()
