"""
Document Parser Module for Voice Studio.

Provides unified document parsing for PPTX, PDF, and TXT input files into
structured `SlideData` objects containing titles, body text, speaker notes,
and raw presentation text.
"""

from dataclasses import dataclass, field, asdict
import os
from pathlib import Path
import re
from typing import List, Dict, Any, Optional, Union

# Try PPTX parser
try:
    import pptx
    _HAS_PPTX = True
except ImportError:
    _HAS_PPTX = False

# Try PyMuPDF / fitz
try:
    import pymupdf as fitz
    _HAS_FITZ = True
except ImportError:
    try:
        import fitz
        _HAS_FITZ = True
    except ImportError:
        _HAS_FITZ = False

# Try pypdf
try:
    import pypdf
    _HAS_PYPDF = True
except ImportError:
    _HAS_PYPDF = False

# Try python-docx
try:
    import docx
    _HAS_DOCX = True
except ImportError:
    _HAS_DOCX = False



@dataclass
class SlideData:
    """Represents text and speaker notes extracted for a single slide page."""
    slide_index: int
    title: str
    body_text: str
    speaker_notes: str
    raw_text: str

    def get_narrative_text(self) -> str:
        """
        Return the primary text intended for voice synthesis.
        Speaker notes take priority if available, otherwise body text.
        """
        raw = ""
        if self.speaker_notes and len(self.speaker_notes.strip()) > 0:
            raw = self.speaker_notes.strip()
        elif self.title and self.body_text:
            raw = f"{self.title}. {self.body_text}".strip()
        else:
            raw = (self.title or self.body_text or self.raw_text).strip()
        return re.sub(r'[\x00-\x09\x0b\x0c\x0e-\x1f\x7f]', ' ', raw).strip()

    def to_dict(self) -> Dict[str, Any]:
        return asdict(self)


@dataclass
class ParsedDocument:
    """Complete parsed document structure holding list of slides and metadata."""
    file_path: str
    doc_type: str  # pptx, pdf, txt, docx
    total_slides: int
    slides: List[SlideData] = field(default_factory=list)

    @property
    def full_text(self) -> str:
        """Combine narrative text from all slides into a single document string."""
        texts = []
        for s in self.slides:
            t = s.get_narrative_text()
            if t:
                texts.append(t)
        return "\n\n".join(texts)

    def to_dict(self) -> Dict[str, Any]:
        return {
            "file_path": str(self.file_path),
            "doc_type": self.doc_type,
            "total_slides": self.total_slides,
            "full_text": self.full_text,
            "slides": [s.to_dict() for s in self.slides],
        }



class DocumentParser:
    """Unified Document Parser for PPTX, PDF, and TXT formats."""

    @staticmethod
    def parse(file_path: Union[str, Path]) -> ParsedDocument:
        """
        Parse presentation or text document file based on extension.
        Supports `.pptx`, `.pdf`, `.txt`, `.md`, `.docx`.
        """
        file_path = Path(file_path)
        if not file_path.exists():
            raise FileNotFoundError(f"Document file not found: {file_path}")

        ext = file_path.suffix.lower()
        if ext == ".pptx":
            return DocumentParser.parse_pptx(file_path)
        elif ext == ".pdf":
            return DocumentParser.parse_pdf(file_path)
        elif ext in (".txt", ".md"):
            return DocumentParser.parse_txt(file_path)
        elif ext == ".docx":
            return DocumentParser.parse_docx(file_path)
        else:
            raise ValueError(f"Unsupported document format: '{ext}'. Supported formats: .pptx, .pdf, .txt, .md, .docx")

    @staticmethod
    def parse_pptx(file_path: Union[str, Path]) -> ParsedDocument:
        """Extract slide titles, body text frames, and speaker notes from PPTX."""
        file_path = Path(file_path)
        if not _HAS_PPTX:
            raise ImportError("python-pptx is required to parse PPTX files. Run `pip install python-pptx`.")

        prs = pptx.Presentation(str(file_path))
        slides_data: List[SlideData] = []

        for idx, slide in enumerate(prs.slides, start=1):
            title = ""
            body_parts = []
            notes_text = ""

            # 1. Extract shapes / text frames
            for shape in slide.shapes:
                if not shape.has_text_frame:
                    continue
                tf = shape.text_frame
                text = tf.text.strip()
                if not text:
                    continue

                is_title = False
                try:
                    if shape == slide.shapes.title:
                        is_title = True
                    elif getattr(shape, "is_placeholder", False):
                        if shape.placeholder_format and shape.placeholder_format.idx == 0:
                            is_title = True
                except Exception:
                    pass

                if is_title and not title:
                    title = text
                else:
                    body_parts.append(text)

            body_text = "\n".join(body_parts)

            # 2. Extract Speaker Notes
            if slide.has_notes_slide and slide.notes_slide:
                notes_tf = slide.notes_slide.notes_text_frame
                if notes_tf:
                    notes_text = notes_tf.text.strip()

            raw_text = f"{title}\n{body_text}".strip()

            slides_data.append(
                SlideData(
                    slide_index=idx,
                    title=title or f"Slide {idx}",
                    body_text=body_text,
                    speaker_notes=notes_text,
                    raw_text=raw_text,
                )
            )

        return ParsedDocument(
            file_path=str(file_path),
            doc_type="pptx",
            total_slides=len(slides_data),
            slides=slides_data,
        )

    @staticmethod
    def parse_pdf(file_path: Union[str, Path]) -> ParsedDocument:
        """Extract text from PDF file page-by-page using PyMuPDF (fitz) or pypdf."""
        file_path = Path(file_path)
        slides_data: List[SlideData] = []

        # 1. Try PyMuPDF (fitz)
        if _HAS_FITZ:
            doc = fitz.open(str(file_path))
            for idx, page in enumerate(doc, start=1):
                text = page.get_text("text").strip()
                lines = [l.strip() for l in text.splitlines() if l.strip()]
                title = lines[0] if lines else f"Page {idx}"
                body_text = "\n".join(lines[1:]) if len(lines) > 1 else ""

                slides_data.append(
                    SlideData(
                        slide_index=idx,
                        title=title,
                        body_text=body_text,
                        speaker_notes="",
                        raw_text=text,
                    )
                )
            doc.close()
            return ParsedDocument(
                file_path=str(file_path),
                doc_type="pdf",
                total_slides=len(slides_data),
                slides=slides_data,
            )

        # 2. Fallback to pypdf
        if _HAS_PYPDF:
            reader = pypdf.PdfReader(str(file_path))
            for idx, page in enumerate(reader.pages, start=1):
                text = page.extract_text() or ""
                text = text.strip()
                lines = [l.strip() for l in text.splitlines() if l.strip()]
                title = lines[0] if lines else f"Page {idx}"
                body_text = "\n".join(lines[1:]) if len(lines) > 1 else ""

                slides_data.append(
                    SlideData(
                        slide_index=idx,
                        title=title,
                        body_text=body_text,
                        speaker_notes="",
                        raw_text=text,
                    )
                )
            return ParsedDocument(
                file_path=str(file_path),
                doc_type="pdf",
                total_slides=len(slides_data),
                slides=slides_data,
            )

        raise ImportError("Either PyMuPDF (fitz) or pypdf is required to parse PDF documents.")

    @staticmethod
    def parse_txt(file_path: Union[str, Path]) -> ParsedDocument:
        """Parse plain text / markdown file into logical slide structures."""
        file_path = Path(file_path)
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            content = f.read().strip()

        # Split by section headers (# or --- or double empty lines)
        if re.search(r"^#+\s+", content, flags=re.MULTILINE):
            sections = re.split(r"(?m)^(?=#+\s+)", content)
        elif "---" in content:
            sections = content.split("---")
        else:
            sections = content.split("\n\n\n")

        slides_data: List[SlideData] = []
        for idx, sec in enumerate(sections, start=1):
            sec_clean = sec.strip()
            if not sec_clean:
                continue
            # Drop markdown horizontal rules / frontmatter fences so the
            # TTS never speaks "dash dash dash" (---, ***, ___ lines).
            sec_clean = re.sub(r"(?m)^\s*[-*_]{3,}\s*$", "", sec_clean).strip()
            if not sec_clean:
                continue

            lines = [l.strip() for l in sec_clean.splitlines() if l.strip()]
            title_line = lines[0] if lines else f"Slide {idx}"
            # Clean markdown header prefixes
            title = re.sub(r"^#+\s*", "", title_line).strip()
            body_text = "\n".join(lines[1:]) if len(lines) > 1 else ""

            slides_data.append(
                SlideData(
                    slide_index=idx,
                    title=title,
                    body_text=body_text,
                    speaker_notes="",
                    raw_text=sec_clean,
                )
            )

        if not slides_data:
            slides_data.append(
                SlideData(
                    slide_index=1,
                    title="Slide 1",
                    body_text=content,
                    speaker_notes="",
                    raw_text=content,
                )
            )

        return ParsedDocument(
            file_path=str(file_path),
            doc_type="txt",
            total_slides=len(slides_data),
            slides=slides_data,
        )

    @staticmethod
    def parse_docx(file_path: Union[str, Path]) -> ParsedDocument:
        """Extract text sections/pages from DOCX files using python-docx or zipfile XML fallback."""
        file_path = Path(file_path)
        paragraphs_text: List[str] = []

        if _HAS_DOCX:
            try:
                doc = docx.Document(str(file_path))
                for p in doc.paragraphs:
                    t = p.text.strip()
                    if t:
                        paragraphs_text.append(t)
                for table in doc.tables:
                    for row in table.rows:
                        row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                        if row_text:
                            paragraphs_text.append(row_text)
            except Exception:
                paragraphs_text = []

        if not paragraphs_text:
            # Fallback using zipfile and xml parsing
            import zipfile
            import xml.etree.ElementTree as ET
            try:
                with zipfile.ZipFile(str(file_path)) as z:
                    xml_content = z.read("word/document.xml")
                    root = ET.fromstring(xml_content)
                    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
                    for p in root.findall(".//w:p", ns):
                        texts = [t.text for t in p.findall(".//w:t", ns) if t.text]
                        p_str = "".join(texts).strip()
                        if p_str:
                            paragraphs_text.append(p_str)
            except Exception as exc:
                raise RuntimeError(f"Failed to parse DOCX file {file_path}: {exc}") from exc

        slides_data: List[SlideData] = []
        chunk_size = 4
        chunks = [paragraphs_text[i:i + chunk_size] for i in range(0, len(paragraphs_text), chunk_size)]

        for idx, chunk in enumerate(chunks, start=1):
            title = chunk[0] if chunk else f"Section {idx}"
            body_text = "\n".join(chunk[1:]) if len(chunk) > 1 else chunk[0]
            raw_text = "\n".join(chunk)
            slides_data.append(
                SlideData(
                    slide_index=idx,
                    title=title,
                    body_text=body_text,
                    speaker_notes="",
                    raw_text=raw_text,
                )
            )

        if not slides_data:
            slides_data.append(
                SlideData(
                    slide_index=1,
                    title="Doc 1",
                    body_text="Empty DOCX document",
                    speaker_notes="",
                    raw_text="",
                )
            )

        return ParsedDocument(
            file_path=str(file_path),
            doc_type="docx",
            total_slides=len(slides_data),
            slides=slides_data,
        )

